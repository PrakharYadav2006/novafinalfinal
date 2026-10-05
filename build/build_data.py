import sys, json, re
sys.path.insert(0,'/tmp/claude-0/-home-claude/84c6249e-43fc-52c8-9191-239551ec3fc2/scratchpad')
import compute2 as C
import openpyxl

# --- final scorecard (from the verified v2 workbook) ---
wb=openpyxl.load_workbook('/home/claude/NovaDrive_Supplier_Risk_Scorecard_v2.xlsx',data_only=True)
calc=wb['Calc']; hdr=[c.value for c in calc[1]]
CALC={}
for r in calc.iter_rows(min_row=2,values_only=True):
    if r[0] is None or not str(r[0]).startswith('ORG'): continue
    CALC[r[0]]=dict(zip(hdr,r))
sc=wb['Scorecard']; SCH=[c.value for c in sc[4]]
SCORE={}
for r in sc.iter_rows(min_row=5,values_only=True):
    if r[1] is None or not str(r[1]).startswith('ORG'): continue
    SCORE[r[1]]=dict(zip(SCH,r))
assert len(CALC)==24 and len(SCORE)==24

# --- tiering workbook: link rows with quotes ---
wbT=openpyxl.load_workbook('/home/claude/NovaDrive_Supplier_Tiering.xlsx',data_only=True)
nl=wbT['Network Links']; nh=[c.value for c in nl[1]]
LINKROWS=[]
for r in nl.iter_rows(min_row=2,values_only=True):
    if r[0] is None: continue
    d=dict(zip(nh,r))
    seller=re.search(r"\((ORG-\d+)\)",d['Seller (legal entity)']).group(1)
    m=re.search(r"\((ORG-\d+)\)",d['Buyer'])
    buyer=m.group(1) if m else 'NOVADRIVE'
    LINKROWS.append(dict(id=d['Link ID'],tier=d['Seller tier'],seller=seller,buyer=buyer,
        comps=d['Component / program'],prods=d['Products reached'],site=d['Seller facility (zone)'],
        status=d['Status'],direct=d['Direct evidence IDs'],context=d['Context evidence IDs'],
        fams=d['# independent source families (direct)'],famids=d['Source family IDs'],dates=d['Evidence dates'],
        quotes=d['Verbatim quotes'],why=d['Reasoning'],flag=d['Flag']))
assert len(LINKROWS)==32

def r1(x): return None if x is None else round(float(x)+1e-9,1)
def r2(x): return None if x is None else round(float(x),2)

COMPNAME={k:v['Component Name'] for k,v in C.COMPS.items()}
PNAME={'P1':'DriveCore','P2':'ChargeBridge','P3':'StoreLink'}
SHORT={k:re.sub(r' Ltd\.$','',v['Legal Name']) for k,v in C.UNI.items()}

sup={}
for n in C.NODES:
    c=CALC[n]; s=SCORE[n]; sc_=C.scores(n); st=C.ST[n]; ind=C.IND[n]
    # cross-check against the verified workbook
    for k,key in [('d1','D1 Financial & operational'),('d2','D2 Location & systemic'),('d3','D3 Network dependence'),('d4','D4 Criticality & substitutability'),('d5','D5 External threats'),('d6','D6 Information gap')]:
        assert abs(sc_[k]-c[key])<1e-6,(n,k,sc_[k],c[key])
    ov=C.overall(sc_); assert abs(ov-c['OVERALL RISK % (active)'])<1e-6
    comps=sorted({x for l in C.LINKS if l['seller']==n for x in l['comps']})
    sells=[]
    for l in C.LINKS:
        if l['seller']==n:
            b='NovaDrive' if l['buyer']=='NOVADRIVE' else SHORT[l['buyer']]
            sells.append(dict(to=b,comps=l['comps'],link=l['id'],status=l['status']))
    buys=[dict(frm=SHORT[l['seller']],id=l['seller'],link=l['id'],comps=l['comps'],status=l['status']) for l in C.LINKS if l['buyer']==n]
    # sub-scores
    subs=[]
    if ind['Current Ratio'] is not None and ind['Net Debt / EBITDA'] is not None:
        subs.append(('Current ratio',ind['Current Ratio'],round(C.lo_(ind['Current Ratio'],*C.SC['cr']),1),'lower is riskier'))
        subs.append(('Net debt / EBITDA',ind['Net Debt / EBITDA'],round(C.hi(ind['Net Debt / EBITDA'],*C.SC['nd']),1),'higher is riskier'))
    else:
        subs.append(('Current ratio',None,None,'not provided'))
        subs.append(('Net debt / EBITDA',None,None,'not provided'))
    subs.append(('On-time shipments, 3-month avg (%)',ind['Shipment Timeliness - 3M Avg (%)'],round(C.lo_(ind['Shipment Timeliness - 3M Avg (%)'],*C.SC['tl']),1),'lower is riskier'))
    subs.append(('On-time change, Jan to latest (pts)',ind['Timeliness Change - Jan To Latest (pp)'],round(C.lo_(ind['Timeliness Change - Jan To Latest (pp)'],*C.SC['tc']),1),'more negative is riskier'))
    ev=[]
    if st and C.UNI[n]['Facility Zone']=='Z01': ev.append(dict(id='EV-001',sev=40))
    if n==C.NAME2ID['Meridian Dielectrics Ltd.']: ev.append(dict(id='EV-003',sev=60))
    sup[n]=dict(
        id=n,name=C.UNI[n]['Legal Name'],short=SHORT[n],tier=c['Tier'],tiern=int(c['Tier'][-1]),zone=c['Facility zone'],
        site=C.UNI[n]['Primary Facility ID'],status=c['Link status'],
        rank=c['Rank (active)'],rankA=c['Rank Test A'],rankB=c['Rank Test B'],
        overall=r1(c['OVERALL RISK % (active)']),overall_raw=round(c['OVERALL RISK % (active)'],2),band=c['Band'],
        conf=c['Confidence'],conf_ev=c['Evidence confidence'],conf_data=c['Data completeness'],
        d=[r1(sc_[k]) for k in ('d1','d2','d3','d4','d5','d6')],
        driver=s['Key vulnerability driver'],next=c['Next step'],
        comps=comps,prods=st['prods'],rev=st['rev'],revpct=round(st['rev_pct']),weeks=st['weeks'],
        exposed=st['exposed'],t1=st['t1_aff'],single=st['single'],fams=C.FAMS[n],
        subs=subs,missing=sc_['missing'],
        idx=[ind['Physical Hazard Index'],ind['Logistics Friction Index'],ind['Infrastructure Index']],
        ass=ind['Assurance Gap Index'],
        sells=sells,buys=buys,events=ev,
        cap=C.UNI[n]['Public Capability'],
        sole=C.SOLE_EXPLICIT.get(n),
        missing_fields=[f for f in C.FIELDS if ind[f] is None],
        single_why=(('Explicit statement that no second source is qualified (%s)'%C.SOLE_EXPLICIT[n]) if n in C.SOLE_EXPLICIT else
                    ('Sole disclosed Tier-1 for its component at 100% of NovaDrive need') if (n in C.ALLOC and all(v==100 for v in C.ALLOC[n].values())) else
                    ('Aster and Boreal share one parent, so they count as one alternative (DOC-075)') if n in C.PARENT_PAIR else
                    ('A second independent Tier-1 is disclosed for the same component') if n in C.ALLOC else
                    ('No alternative is stated in the data. Not proof that none exists')),
        open_note=({'ref':'U1','text':'Customer index mentions M10/M20 applications but does not identify the receiving line or purchase allocation.','settle':'Direct IonPeak disclosure, or a shipment record naming Verdant and IonPeak with component scope.'} if C.STATUS[n]=='Inferred' else None),
        parent_pair=(n in C.PARENT_PAIR),
    )
# zone indices
zi={}
for e,u in C.UNI.items():
    z=u['Facility Zone']; i=C.IND.get(e)
    if i and i['Physical Hazard Index'] is not None:
        zi.setdefault(z,set()).add((i['Physical Hazard Index'],i['Logistics Friction Index'],i['Infrastructure Index']))
print({z:sorted(v) for z,v in zi.items()})
ZIDX={}
for z,v in zi.items():
    if len(v)==1: ZIDX[z]=list(next(iter(v)))
    else: print('ZONE INDEX NOT UNIQUE',z,v)
# evidence records
EV={}
for eid,e in C.EVD.items():
    EV[eid]=dict(date=str(e['Evidence Date'])[:10],type=e['Evidence Type'],fam=e['Source Family ID'],title=e['Title / Parties'],detail=e['Evidence Detail'],status=e['Record Status'])
# events
ews=wb['Events']; eh=[c.value for c in ews[1]]
EVENTS=[dict(zip(eh,r)) for r in ews.iter_rows(min_row=2,values_only=True) if r[0] and str(r[0]).startswith('EV')]
raw_ev={r['Event ID']:r for r in C.rows('Risk Event Flags')}
EVJ=[dict(id=e['Event'],date=str(e['Date'])[:10],head=e['Headline'],decision=e['Decision'],matched=e['Matched to'],sev=e['Severity (0-100)'],why=e['Why'],verify=e['What to verify'],detail=raw_ev[e['Event']]['Event Detail'],fam=raw_ev[e['Event']]['Source Family ID']) for e in EVENTS]
# non-network entities by zone (from Rejected/Unresolved sheets)
NONNET={
 'Z02':[dict(name='Delta Consumer Plastics Ltd.',ref='R2',why='Name looks like Delta Capacitor Works but it is a different company and never became a supply award.')],
 'Z05':[dict(name='CommonSpan Holdings Ltd.',ref='R7 / S2',why='Holding company. Shown as the common parent of Aster and Boreal, not as a supplier tier.'),
        dict(name='IonPeak Western Development Center (SITE-900)',ref='R4',why='Research and pilot site on a different process. Not a qualified die source.')],
 'Z07':[dict(name='Harbor Freight Services Ltd.',ref='R5',why='Freight consolidator, not a manufacturer.'),
        dict(name='Ion Peak Trading Ltd.',ref='R3',why='Unrelated broker, not IonPeak Semiconductor.'),
        dict(name='Jade Circuits Holdings Ltd.',ref='R8',why='Jade parent and head office. Does not make substrates.')],
 'Z08':[dict(name='Solace Optics Ltd.',ref='U2',why='Hypothesis only (sample evaluation). Not placed in the network.')],
}
meta=dict(comp=COMPNAME,pname=PNAME,weights=dict(zip(['D1','D2','D3','D4','D5','D6'],[20,15,30,15,10,10])),
          band_cuts=[41.5,54.3],rev={'P1':624,'P2':520,'P3':416},total=1560)
out=dict(sup=sup,links=LINKROWS,ev=EV,events=EVJ,zidx=ZIDX,nonnet=NONNET,meta=meta)
json.dump(out,open('data.json','w'),default=str,separators=(',',':'))
import os; print(os.path.getsize('data.json'))
for n in sorted(sup,key=lambda k:sup[k]['rank']): s=sup[n]; print(s['rank'],s['short'],s['zone'],s['tier'],s['overall'],s['band'],s['conf'],s['comps'])
